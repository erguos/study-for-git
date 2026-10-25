"""Ingest uploaded files (CSV/Excel tables and text documents) into MySQL."""
import io
import re
from typing import List, Optional

import pandas as pd
from sqlalchemy import text
from sqlalchemy.orm import Session

from ..models import Dataset, DocumentChunk

MAX_CHUNK_LEN = 400


def _mysql_type(dtype) -> str:
    name = str(dtype)
    if name.startswith("int"):
        return "BIGINT"
    if name.startswith("float") or name.startswith("double"):
        return "DOUBLE"
    if name.startswith("bool"):
        return "TINYINT(1)"
    try:
        import datetime
        if issubclass(dtype.type, datetime.datetime):
            return "DATETIME"
    except Exception:
        pass
    return "TEXT"


def _clean_col(name: str) -> str:
    name = re.sub(r"[^0-9a-zA-Z_\u4e00-\u9fa5]", "_", str(name)).strip("_")
    if not name or name[0].isdigit():
        name = "c_" + name
    return name


def ingest_table(
    db: Session,
    file_bytes: bytes,
    file_name: str,
    description: str,
) -> dict:
    """Parse CSV/Excel, create a MySQL table and return column metadata."""
    ext = file_name.rsplit(".", 1)[-1].lower()
    if ext == "csv":
        df = pd.read_csv(io.BytesIO(file_bytes))
    elif ext in ("xlsx", "xls"):
        df = pd.read_excel(io.BytesIO(file_bytes))
    else:
        raise ValueError("不支持的表格格式，请上传 CSV 或 Excel 文件")

    df = df.dropna(how="all").reset_index(drop=True)
    if df.empty or df.shape[1] == 0:
        raise ValueError("文件内容为空")

    cols = [{"name": str(c), "clean": _clean_col(str(c)), "dtype": df[c].dtype}
            for c in df.columns]
    # avoid duplicate cleaned names
    seen = {}
    for c in cols:
        base = c["clean"]
        n = seen.get(base, 0)
        seen[base] = n + 1
        c["clean"] = base if n == 0 else f"{base}_{n}"

    dataset = Dataset(
        name=file_name,
        type="table",
        description=description,
        status="creating",
    )
    db.add(dataset)
    db.flush()

    table_name = f"dataset_{dataset.id}"
    dataset.table_name = table_name

    col_defs = ", ".join(
        f"`{c['clean']}` {_mysql_type(c['dtype'])} NULL" for c in cols
    )
    create_sql = f"CREATE TABLE `{table_name}` (id INT AUTO_INCREMENT PRIMARY KEY, {col_defs})"
    db.execute(text(create_sql))

    columns_clean = [c["clean"] for c in cols]
    placeholders = ", ".join(["%s"] * len(columns_clean))
    insert_sql = f"INSERT INTO `{table_name}` ({', '.join('`' + c + '`' for c in columns_clean)}) VALUES ({placeholders})"
    data = df[list(c["name"] for c in cols)].where(pd.notnull(df), None).values.tolist()
    data = [tuple(row) for row in data]
    db.connection().exec_driver_sql(insert_sql, data)

    schema = [
        {"name": c["clean"], "original": c["name"], "type": _mysql_type(c["dtype"])}
        for c in cols
    ]
    dataset.schema_json = schema
    dataset.status = "ready"
    db.commit()
    db.refresh(dataset)

    return {
        "dataset_id": dataset.id,
        "name": dataset.name,
        "type": "table",
        "rows": len(data),
        "columns": schema,
        "table_name": table_name,
        "message": f"已导入 {len(data)} 行，{len(schema)} 列",
    }


def _extract_text(ext: str, file_bytes: bytes) -> str:
    if ext in ("txt", "md", "text"):
        try:
            return file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            return file_bytes.decode("gbk", errors="ignore")
    if ext in ("csv",):
        return file_bytes.decode("utf-8", errors="ignore")
    if ext == "docx":
        try:
            import docx  # python-docx
            doc = docx.Document(io.BytesIO(file_bytes))
            return "\n".join(p.text for p in doc.paragraphs)
        except ImportError:
            raise ValueError("解析 docx 需要安装 python-docx")
    if ext == "pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(file_bytes))
            return "\n".join((page.extract_text() or "") for page in reader.pages)
        except ImportError:
            raise ValueError("解析 pdf 需要安装 pypdf")
    raise ValueError(f"不支持的文档格式: {ext}")


def _split_chunks(text: str) -> List[str]:
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return []
    return [text[i:i + MAX_CHUNK_LEN] for i in range(0, len(text), MAX_CHUNK_LEN)]


def ingest_doc(
    db: Session,
    file_bytes: bytes,
    file_name: str,
    description: str,
) -> dict:
    ext = file_name.rsplit(".", 1)[-1].lower()
    text = _extract_text(ext, file_bytes)
    chunks = _split_chunks(text)

    dataset = Dataset(
        name=file_name,
        type="doc",
        description=description,
        status="ready",
        chunk_count=len(chunks),
    )
    db.add(dataset)
    db.flush()

    for i, chunk in enumerate(chunks):
        db.add(DocumentChunk(
            dataset_id=dataset.id,
            chunk_index=i,
            content=chunk,
            source_file=file_name,
        ))
    db.commit()
    db.refresh(dataset)
    return {
        "dataset_id": dataset.id,
        "name": dataset.name,
        "type": "doc",
        "chunks": len(chunks),
        "message": f"已解析 {file_name}，切分为 {len(chunks)} 个片段",
    }
